"""Điền dữ liệu vào mẫu DOCX, bao gồm template dùng nhãn sẵn có và {{biến}}."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.document import Document as DocumentType
from docx.text.paragraph import Paragraph

from services.extraction_service import FIELD_NAMES
from services.template_service import FormTemplate
from utils.text_utils import normalize_citizen_id, normalize_date, normalize_phone, normalize_text


class DocumentGenerationError(RuntimeError):
    pass


def _set_field_value(value: str | None) -> str:
    return normalize_text(value) or ""


def _append_after_label(paragraphs: list[Paragraph], label: str, value: str) -> None:
    """Điền giá trị vào vùng trống sau nhãn của một dòng một trường."""
    if not value:
        return
    for paragraph in paragraphs:
        if label.lower() in paragraph.text.lower():
            _replace_after_label(paragraph, label, value)
            return


def _replace_after_label(paragraph: Paragraph, label: str, value: str) -> bool:
    """Thay phần trống sau nhãn, kể cả khi mẫu dùng tab leader dấu chấm."""
    if not value:
        return False
    text = paragraph.text
    folded = text.casefold()
    label_start = folded.find(label.casefold())
    if label_start < 0:
        return False

    start = label_start + len(label)
    if not label.rstrip().endswith(":"):
        colon = text.find(":", start)
        if colon >= 0:
            start = colon + 1
    return _replace_text_span(paragraph, start, len(text), f" {value}")


def _find_labeled_paragraph(paragraphs: list[Paragraph], label: str) -> Paragraph | None:
    label_lower = label.lower()
    return next((paragraph for paragraph in paragraphs if label_lower in paragraph.text.lower()), None)


def _replace_text_span(paragraph: Paragraph, start: int, end: int, value: str) -> bool:
    """Thay một đoạn văn bản, kể cả khi Word đã chia nó thành nhiều run.

    Không gán ``paragraph.text`` vì cách đó làm mất tab, chấm dẫn và định dạng
    nghiêng trong biểu mẫu gốc.
    """
    if start < 0 or end < start:
        return False

    offset = 0
    inserted = False
    for run in paragraph.runs:
        original = run.text
        run_start = offset
        run_end = offset + len(original)
        offset = run_end

        if run_end <= start or run_start >= end:
            continue

        remove_start = max(start, run_start) - run_start
        remove_end = min(end, run_end) - run_start
        before = original[:remove_start]
        after = original[remove_end:]
        if not inserted:
            run.text = f"{before}{value}{after}"
            inserted = True
        else:
            run.text = f"{before}{after}"
    return inserted


def _replace_between_labels(paragraph: Paragraph, left_label: str, right_label: str, value: str) -> bool:
    """Điền vào vùng trống nằm giữa hai nhãn trên cùng một dòng."""
    if not value:
        return False
    text = paragraph.text
    folded = text.casefold()
    left = folded.find(left_label.casefold())
    if left < 0:
        return False
    start = left + len(left_label)
    right = folded.find(right_label.casefold(), start)
    if right < 0:
        return False
    return _replace_text_span(paragraph, start, right, f" {value} ")


def _fill_identity_line(paragraphs: list[Paragraph], fields: dict[str, str]) -> None:
    """Điền ba ô Ngày sinh, Giới tính, Dân tộc của dòng số 2.

    Ba ô này vốn ở chung một paragraph.  Chúng không thể dùng cách thêm cuối
    dòng, vì dữ liệu sẽ vượt qua cả nhãn Giới tính và Dân tộc.
    """
    paragraph = _find_labeled_paragraph(paragraphs, "2. Ngày, tháng, năm sinh")
    if paragraph is None:
        return
    _replace_between_labels(paragraph, "Ngày, tháng, năm sinh:", "Giới tính:", fields.get("date_of_birth", ""))
    _replace_between_labels(paragraph, "Giới tính:", "Dân tộc:", fields.get("gender", ""))
    _replace_after_label(paragraph, "Dân tộc:", fields.get("ethnic_group", ""))


def _fill_bank_account_line(paragraphs: list[Paragraph], fields: dict[str, str]) -> None:
    """Điền riêng số tài khoản và tên ngân hàng trên dòng số 10."""
    paragraph = _find_labeled_paragraph(paragraphs, "- Số tài khoản")
    if paragraph is None:
        return
    _replace_between_labels(
        paragraph,
        "Số tài khoản:",
        "Ngân hàng:",
        fields.get("bank_account_number", ""),
    )
    _replace_after_label(paragraph, "Ngân hàng:", fields.get("bank_name", ""))


def _replace_placeholders(paragraph: Paragraph, fields: dict[str, str]) -> None:
    # Mẫu mới nên đặt {{full_name}} trong một run. Trường hợp nhiều run vẫn có
    # fallback thay bằng text toàn đoạn để template tiếp tục hoạt động.
    for run in paragraph.runs:
        for key, value in fields.items():
            token = f"{{{{{key}}}}}"
            if token in run.text:
                run.text = run.text.replace(token, value)
    text = paragraph.text
    if any(f"{{{{{key}}}}}" in text for key in fields):
        for key, value in fields.items():
            text = text.replace(f"{{{{{key}}}}}", value)
        paragraph.text = text


def _all_paragraphs(document: DocumentType) -> list[Paragraph]:
    paragraphs = list(document.paragraphs)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)
    for section in document.sections:
        paragraphs.extend(section.header.paragraphs)
        paragraphs.extend(section.footer.paragraphs)
    return paragraphs


def _fill_retirement_assistance_template(document: DocumentType, fields: dict[str, str]) -> None:
    """Mapping nhãn của mẫu trợ cấp hưu trí hiện có trong thư mục templates_word."""
    paragraphs = _all_paragraphs(document)
    full_name = fields.get("full_name", "")

    _append_after_label(paragraphs, "1. Họ, chữ đệm, tên", full_name)
    _fill_identity_line(paragraphs, fields)
    _append_after_label(paragraphs, "3. Thẻ Căn cước hoặc số định danh cá nhân", fields.get("citizen_id", ""))
    _append_after_label(paragraphs, "4. Nơi cư trú", fields.get("residence_address", ""))
    _append_after_label(paragraphs, "5. Địa chỉ liên lạc", fields.get("contact_address", ""))
    _append_after_label(paragraphs, "6. Số điện thoại", fields.get("phone_number", ""))
    _append_after_label(paragraphs, "9. Nơi đề nghị nhận trợ cấp", fields.get("benefit_receiving_location", ""))
    _append_after_label(paragraphs, "- Tên tài khoản", fields.get("bank_account_name", ""))
    _fill_bank_account_line(paragraphs, fields)
    _append_after_label(paragraphs, "II. Thông tin người giám hộ", "")
    # Các nhãn ở mục II trùng mục I, nên tìm theo chỉ mục từ vị trí tiêu đề.
    guardian_start = next((index for index, p in enumerate(paragraphs) if "II. Thông tin người giám hộ" in p.text), len(paragraphs))
    guardian_paragraphs = paragraphs[guardian_start:]
    _append_after_label(guardian_paragraphs, "1. Họ, chữ đệm, tên", fields.get("guardian_full_name", ""))
    _append_after_label(guardian_paragraphs, "2. Ngày, tháng, năm sinh", fields.get("guardian_date_of_birth", ""))
    _append_after_label(guardian_paragraphs, "3. Thẻ Căn cước hoặc số định danh cá nhân", fields.get("guardian_citizen_id", ""))
    _append_after_label(guardian_paragraphs, "4. Địa chỉ liên hệ", fields.get("guardian_address", ""))
    _append_after_label(guardian_paragraphs, "5. Số điện thoại", fields.get("guardian_phone", ""))
    _append_after_label(guardian_paragraphs, "6. Quan hệ với người đề nghị", fields.get("guardian_relationship", ""))


def generate_document(template: FormTemplate, raw_fields: dict, output_path: str | Path) -> Path:
    if not template.path.is_file():
        raise DocumentGenerationError(f"Không tìm thấy mẫu Word: {template.file_name}")
    # Luôn có sẵn mọi trường đã biết (mặc định rỗng) trước khi ghi đè bằng
    # raw_fields: nếu thiếu, {{placeholder}} tương ứng trong mẫu mới sẽ bị bỏ
    # sót và hiện nguyên văn "{{...}}" trong tệp Word thay vì để trống.
    fields = {key: "" for key in FIELD_NAMES}
    fields.update({key: _set_field_value(value) for key, value in raw_fields.items()})
    fields["citizen_id"] = normalize_citizen_id(fields.get("citizen_id"))
    fields["date_of_birth"] = normalize_date(fields.get("date_of_birth"))
    fields["guardian_date_of_birth"] = normalize_date(fields.get("guardian_date_of_birth"))
    fields["phone_number"] = normalize_phone(fields.get("phone_number"))
    fields["guardian_phone"] = normalize_phone(fields.get("guardian_phone"))
    fields["deceased_citizen_id"] = normalize_citizen_id(fields.get("deceased_citizen_id"))
    fields["deceased_date_of_birth"] = normalize_date(fields.get("deceased_date_of_birth"))
    fields["deceased_death_date"] = normalize_date(fields.get("deceased_death_date"))
    fields["death_certificate_date"] = normalize_date(fields.get("death_certificate_date"))
    fields["org_phone"] = normalize_phone(fields.get("org_phone"))

    try:
        document = Document(template.path)
        paragraphs = _all_paragraphs(document)
        for paragraph in paragraphs:
            _replace_placeholders(paragraph, fields)
        if template.id == "tro_cap_huu_tri":
            _fill_retirement_assistance_template(document, fields)
        target = Path(output_path)
        target.parent.mkdir(parents=True, exist_ok=True)
        document.save(target)
        return target
    except (OSError, ValueError) as error:
        raise DocumentGenerationError(str(error)) from error
