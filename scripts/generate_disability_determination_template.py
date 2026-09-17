"""Tạo mẫu DOCX Đơn đề nghị xác định mức độ khuyết tật (Mẫu số 01, TT 01/2019/TT-BLĐTBXH).

Mục III (dạng khuyết tật, mức độ khuyết tật) được tái tạo đầy đủ để giữ đúng bố
cục của đơn chính thức, nhưng để trống cho Hội đồng xác định khuyết tật/người
dân đánh dấu tay: đây là nội dung do chuyên môn y tế đánh giá trực tiếp, không
phải thông tin có thể trích xuất đáng tin cậy từ OCR hay phỏng vấn giọng nói.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "templates_word" / "xac_dinh_khuyet_tat.docx"

DISABILITY_TYPE_GROUPS = (
    (
        "1",
        "Khuyết tật vận động",
        (
            ("1.1", "Mềm nhẽo hoặc co cứng toàn thân"),
            ("1.2", "Thiếu tay hoặc không cử động được tay"),
            ("1.3", "Thiếu chân hoặc không cử động được chân"),
            ("1.4", "Yếu, liệt, teo cơ hoặc hạn chế vận động tay, chân, lưng, cổ"),
            (
                "1.5",
                "Cong, vẹo, chân tay, lưng, cổ; gù cột sống lưng hoặc dị dạng, "
                "biến dạng khác trên cơ thể ở đầu, cổ, lưng, tay, chân",
            ),
            ("1.6", "Có kết luận của cơ sở y tế cấp tỉnh trở lên về suy giảm chức năng vận động"),
        ),
    ),
    (
        "2",
        "Khuyết tật nghe, nói",
        (
            ("2.1", "Không phát ra âm thanh, lời nói"),
            ("2.2", "Phát ra âm thanh, lời nói nhưng không rõ tiếng, rõ câu"),
            ("2.3", "Không nghe được"),
            ("2.4", "Khiếm khuyết hoặc dị dạng cơ quan phát âm ảnh hưởng đến việc phát âm"),
            ("2.5", "Khiếm khuyết hoặc dị dạng vành tai hoặc ống tai ngoài ảnh hưởng đến nghe"),
            ("2.6", "Có kết luận của cơ sở y tế cấp tỉnh trở lên về suy giảm chức năng nghe, nói"),
        ),
    ),
    (
        "3",
        "Khuyết tật nhìn",
        (
            ("3.1", "Mù một hoặc hai mắt"),
            ("3.2", "Thiếu một hoặc hai mắt"),
            ("3.3", "Khó khăn khi nhìn hoặc không nhìn thấy các đồ vật"),
            ("3.4", "Khó khăn khi phân biệt màu sắc hoặc không phân biệt được các màu sắc"),
            ("3.5", "Rung, giật nhãn thị, đục nhân mắt hoặc sẹo loét giác mạc"),
            ("3.6", "Bị dị tật, biến dạng ở vùng mắt"),
            ("3.7", "Có kết luận của cơ sở y tế cấp tỉnh trở lên về suy giảm chức năng nhìn"),
        ),
    ),
    (
        "4",
        "Khuyết tật thần kinh, tâm thần",
        (
            ("4.1", "Thường ngồi một mình, chơi một mình, không bao giờ nói chuyện hoặc quan tâm tới bất kỳ ai"),
            (
                "4.2",
                "Có những hành vi bất thường như kích động, cáu giận hoặc sợ hãi vô cớ gây ảnh hưởng "
                "đến sức khỏe, sự an toàn của bản thân và người khác",
            ),
            (
                "4.3",
                "Bất ngờ dừng mọi hoạt động, mắt mở trừng trừng không chớp, co giật chân tay, môi, mặt "
                "hoặc bất thình lình ngã xuống, co giật, sùi bọt mép, gọi hỏi không biết",
            ),
            ("4.4", "Bị mất trí nhớ, bỏ nhà đi lang thang"),
            ("4.5", "Có kết luận của cơ sở y tế cấp tỉnh trở lên về suy giảm thần kinh, tâm thần"),
        ),
    ),
    (
        "5",
        "Khuyết tật trí tuệ",
        (
            (
                "5.1",
                "Khó khăn trong việc nhận biết người thân trong gia đình hoặc khó khăn trong giao tiếp "
                "với những người xung quanh so với người cùng lứa tuổi",
            ),
            ("5.2", "Chậm chạp, ngờ nghệch hoặc không thể làm được một việc đơn giản (so với tuổi) dù đã được hướng dẫn"),
            (
                "5.3",
                "Khó khăn trong việc đọc, viết, tính toán và kỹ năng học tập khác so với người cùng tuổi "
                "do chậm phát triển trí tuệ",
            ),
            ("5.4", "Có kết luận cơ sở y tế cấp tỉnh trở lên về chậm phát triển trí tuệ"),
        ),
    ),
    (
        "6",
        "Khuyết tật khác",
        (
            (
                "6.1",
                "Có kết luận của cơ sở y tế cấp tỉnh trở lên về bệnh tê bì, mất cảm giác ở tay, chân hoặc "
                "sự bất thường của cơ thể làm giảm khả năng thực hiện các hoạt động; lao động; đọc, viết, "
                "tính toán và kỹ năng học tập khác; sinh hoạt hoặc giao tiếp",
            ),
            (
                "6.2",
                "Có kết luận của cơ sở y tế cấp tỉnh trở lên về bệnh hô hấp hoặc do bệnh tim mạch hoặc do "
                "rối loạn đại, tiểu tiện mặc dù đã được điều trị liên tục trên 3 tháng, làm giảm khả năng "
                "thực hiện các hoạt động; lao động; đọc, viết, tính toán và kỹ năng học tập khác; sinh hoạt "
                "hoặc giao tiếp",
            ),
            ("6.3", "Có kết luận của cơ sở y tế cấp tỉnh trở lên về rối loạn phổ tự kỷ hoặc các loại bệnh hiếm"),
        ),
    ),
)

ABILITY_LEVEL_ACTIVITIES = (
    "1. Đi lại",
    "2. Ăn, uống",
    "3. Tiểu tiện, đại tiện",
    "4. Vệ sinh cá nhân như đánh răng, rửa mặt, tắm rửa...",
    "5. Mặc, cởi quần áo, giầy dép",
    "6. Nghe và hiểu người khác nói gì",
    "7. Diễn đạt được ý muốn và suy nghĩ của bản thân qua lời nói",
    "8. Làm các việc gia đình như gấp quần áo, quét nhà, rửa bát, nấu cơm phù hợp với độ tuổi; lao động, sản xuất tạo thu nhập",
    "9. Giao tiếp xã hội, hòa nhập cộng đồng phù hợp với độ tuổi",
    "10. Đọc, viết, tính toán và kỹ năng học tập khác",
)

ABILITY_LEVEL_COLUMNS = (
    "Thực hiện được",
    "Thực hiện được nhưng cần trợ giúp",
    "Không thực hiện được",
    "Không xác định được",
)


def _font(run, *, size: float = 12, bold: bool = False, italic: bool = False) -> None:
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def _p(document, text: str = "", *, bold=False, italic=False, center=False, indent=0.0, before=0, after=3, size=12):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.left_indent = Cm(indent)
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    _font(paragraph.add_run(text), size=size, bold=bold, italic=italic)
    return paragraph


def _set_column_widths(table, widths_cm: tuple[float, ...]) -> None:
    table.autofit = False
    for row in table.rows:
        for cell, width in zip(row.cells, widths_cm):
            cell.width = Cm(width)


def build() -> Path:
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(1.5)
    section.left_margin, section.right_margin = Cm(1.75), Cm(1.5)

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(12)

    _p(document, "Mẫu số 01", italic=True, center=True, after=0)
    _p(
        document,
        "(Ban hành kèm theo Thông tư số 01/2019/TT-BLĐTBXH ngày 02 tháng 01 năm 2019)",
        italic=True,
        center=True,
        after=6,
        size=11,
    )
    _p(document, "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM", bold=True, center=True, after=0)
    motto = _p(document, "Độc lập - Tự do - Hạnh phúc", bold=True, center=True, after=0)
    motto.runs[0].underline = True
    _p(document, "---------------", center=True, after=8)

    _p(document, "ĐƠN ĐỀ NGHỊ XÁC ĐỊNH, XÁC ĐỊNH LẠI MỨC ĐỘ KHUYẾT TẬT VÀ", bold=True, center=True, size=13, after=0)
    _p(document, "CẤP, CẤP ĐỔI, CẤP LẠI GIẤY XÁC NHẬN KHUYẾT TẬT", bold=True, center=True, size=13, after=8)

    _p(document, "Kính gửi: Chủ tịch Ủy ban nhân dân phường Minh Phụng, TP. Hồ Chí Minh", after=6)
    _p(document, "Sau khi tìm hiểu quy định về xác định mức độ khuyết tật, tôi đề nghị:", after=4)

    procedure_options = (
        ("{{disability_procedure_new_check}} Xác định mức độ khuyết tật và cấp Giấy xác nhận khuyết tật"),
        ("{{disability_procedure_redetermine_check}} Xác định lại mức độ khuyết tật và cấp Giấy xác nhận khuyết tật"),
        ("{{disability_procedure_reissue_check}} Cấp lại Giấy xác nhận khuyết tật"),
        ("{{disability_procedure_replace_check}} Cấp đổi Giấy xác nhận khuyết tật"),
    )
    for text in procedure_options:
        _p(document, text, indent=0.7, after=2)

    _p(
        document,
        "(Trường hợp cấp đổi Giấy xác nhận khuyết tật thì không phải kê khai thông tin tại Mục III dưới đây).",
        italic=True,
        after=6,
    )
    _p(document, "Cụ thể:", after=4)

    _p(document, "I. Thông tin người được xác định mức độ khuyết tật", bold=True, after=3)
    _p(document, "- Họ và tên: {{full_name}}")
    _p(document, "- Ngày, tháng, năm sinh: {{date_of_birth}}     Giới tính: {{gender}}")
    _p(document, "- Số CMND hoặc căn cước công dân: {{citizen_id}}")
    _p(document, "- Hộ khẩu thường trú: {{residence_address}}")
    _p(document, "- Nơi ở hiện nay: {{contact_address}}", after=6)

    _p(document, "II. Thông tin người đại diện hợp pháp (nếu có)", bold=True, after=3)
    _p(document, "- Họ và tên: {{guardian_full_name}}")
    _p(document, "- Mối quan hệ với người được xác định khuyết tật: {{guardian_relationship}}")
    _p(document, "- Số CMND hoặc căn cước công dân: {{guardian_citizen_id}}")
    _p(document, "- Hộ khẩu thường trú: {{guardian_residence_address}}")
    _p(document, "- Nơi ở hiện nay: {{guardian_current_address}}")
    _p(document, "- Số điện thoại: {{guardian_phone}}")

    page_break = document.add_paragraph()
    page_break.add_run().add_break(WD_BREAK.PAGE)

    _p(document, "III. Thông tin về tình trạng khuyết tật", bold=True, after=4)
    _p(
        document,
        "1. Thông tin về dạng khuyết tật (Đánh dấu x vào ô tương ứng)",
        bold=True,
        after=3,
    )
    _p(
        document,
        "Mục này do Hội đồng xác định mức độ khuyết tật/cơ sở y tế đánh giá trực tiếp, "
        "đề nghị đánh dấu x bằng tay sau khi in đơn.",
        italic=True,
        size=10.5,
        after=4,
    )

    type_table = document.add_table(rows=1, cols=4, style="Table Grid")
    _set_column_widths(type_table, (1.3, 12.7, 1.5, 1.5))
    header_cells = type_table.rows[0].cells
    for cell, text in zip(header_cells, ("STT", "Các dạng khuyết tật", "Có", "Không")):
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _font(paragraph.add_run(text), size=11, bold=True)

    for group_number, group_title, items in DISABILITY_TYPE_GROUPS:
        row = type_table.add_row().cells
        row[0].text = ""
        _font(row[0].paragraphs[0].add_run(group_number), size=11, bold=True)
        row[1].text = ""
        _font(row[1].paragraphs[0].add_run(group_title), size=11, bold=True)
        for cell in row[2:]:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for item_number, item_text in items:
            item_row = type_table.add_row().cells
            item_row[0].text = ""
            _font(item_row[0].paragraphs[0].add_run(item_number), size=10.5)
            item_row[1].text = ""
            _font(item_row[1].paragraphs[0].add_run(item_text), size=10.5)
            for cell in item_row[2:]:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    _p(document, "", after=8)
    _p(
        document,
        "2. Thông tin về mức độ khuyết tật (Trường hợp trẻ em dưới 6 tuổi không phải kê khai)",
        bold=True,
        after=4,
    )

    level_table = document.add_table(rows=1, cols=5, style="Table Grid")
    _set_column_widths(level_table, (7.7, 2.4, 2.4, 2.4, 2.1))
    level_header = level_table.rows[0].cells
    for cell, text in zip(level_header, ("Các hoạt động", *ABILITY_LEVEL_COLUMNS)):
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _font(paragraph.add_run(text), size=10.5, bold=True)

    for activity in ABILITY_LEVEL_ACTIVITIES:
        row = level_table.add_row().cells
        row[0].text = ""
        _font(row[0].paragraphs[0].add_run(activity), size=10.5)
        for cell in row[1:]:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    _p(document, "", after=8)

    signature = document.add_table(rows=1, cols=2)
    signature.autofit = False
    for cell in signature.rows[0].cells:
        cell.width = Cm(8.5)
    right = signature.rows[0].cells[1]
    right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    for index, (text, bold, italic) in enumerate(
        (
            ("………, ngày {{declaration_day}} tháng {{declaration_month}} năm {{declaration_year}}", False, True),
            ("Người viết đơn", True, False),
            ("(Ký và ghi rõ họ tên)", False, True),
        )
    ):
        paragraph = right.paragraphs[0] if index == 0 else right.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(1)
        _font(paragraph.add_run(text), size=11.5, bold=bold, italic=italic)

    document.core_properties.title = (
        "Đơn đề nghị xác định, xác định lại mức độ khuyết tật và cấp, cấp đổi, cấp lại Giấy xác nhận khuyết tật"
    )
    document.core_properties.author = "Auto Form"

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build())
