"""Tạo mẫu mai táng A4 hai trang, bám đúng thứ tự của PDF chính thức."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "templates_word" / "ho_tro_mai_tang.docx"


def _font(run, size=11, bold=False, italic=False):
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def _p(doc, text="", *, bold=False, center=False, indent=0, before=0, after=1, size=11):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Cm(indent)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    r = p.add_run(text)
    _font(r, size=size, bold=bold)
    return p


def _no_borders(table):
    props = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = OxmlElement(f"w:{edge}")
        tag.set(qn("w:val"), "nil")
        borders.append(tag)
    props.append(borders)


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(1.35)
    section.left_margin = section.right_margin = Cm(1.7)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(11)

    _p(doc, "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM", bold=True, center=True, after=0)
    _p(doc, "Độc lập - Tự do - Hạnh phúc", bold=True, center=True, after=0)
    _p(doc, "---------------", center=True, after=4)
    _p(doc, "TỜ KHAI ĐỀ NGHỊ HỖ TRỢ CHI PHÍ MAI TÁNG", bold=True, center=True, size=12, after=4)
    _p(doc, "Kính gửi: Chủ tịch Ủy ban nhân dân phường Minh Phụng", center=True, after=5)

    _p(doc, "1. Thông tin người đề nghị (người khai)", bold=True, after=2)
    _p(doc, "Họ, chữ đệm, tên (Viết chữ in hoa): {{full_name}}")
    _p(doc, "Nơi cư trú: {{residence_address}}")
    _p(doc, "Thẻ Căn cước hoặc số định danh cá nhân: {{citizen_id}}")
    _p(doc, "Quan hệ với người chết: {{deceased_relationship}}")
    _p(doc, "Nội dung đề nghị: {{request_content}}", after=3)

    _p(doc, "2. Thông tin người chết được tổ chức mai táng", bold=True, after=2)
    _p(doc, "Họ, chữ đệm, tên: {{deceased_full_name}}")
    _p(doc, "Ngày, tháng, năm sinh: {{deceased_date_of_birth}}")
    _p(doc, "Giới tính: {{deceased_gender}}     Dân tộc: {{deceased_ethnic_group}}     Quốc tịch: {{deceased_nationality}}")
    _p(doc, "Nơi cư trú: {{deceased_residence_address}}")
    _p(doc, "Thẻ Căn cước hoặc số định danh cá nhân: {{deceased_citizen_id}}")
    _p(doc, "Đã chết vào lúc: {{deceased_death_hour}} giờ {{deceased_death_minute}} phút, ngày {{deceased_death_date}}")
    _p(doc, "Nơi chết: {{deceased_death_place}}")
    _p(doc, "Nguyên nhân chết: {{deceased_death_cause}}")
    _p(doc, "Số Giấy báo tử/Giấy tờ thay thế Giấy báo tử: {{death_certificate_number}} do {{death_certificate_issuer}} cấp ngày {{death_certificate_date}}")

    page_break = doc.add_paragraph()
    page_break.add_run().add_break(WD_BREAK.PAGE)

    _p(doc, "3. Người, tổ chức lo mai táng nhận hỗ trợ chi phí mai táng", bold=True, after=2)
    _p(doc, "3.1. Trường hợp cá nhân, thân nhân đứng ra tổ chức mai táng:", bold=True)
    _p(doc, "Họ và tên: {{full_name}}")
    _p(doc, "Ngày, tháng, năm sinh: {{date_of_birth}}     Nam/Nữ: {{gender}}")
    _p(doc, "Thẻ Căn cước hoặc số định danh cá nhân: {{citizen_id}}")
    _p(doc, "Nơi cư trú: {{residence_address}}")
    _p(doc, "Quan hệ với người chết: {{deceased_relationship}}")
    _p(doc, "Số điện thoại liên hệ: {{phone_number}}", after=2)
    _p(doc, "3.2. Trường hợp cơ quan, tổ chức đứng ra tổ chức mai táng:", bold=True)
    _p(doc, "Tên tổ chức: {{org_name}}")
    _p(doc, "Địa chỉ: {{org_address}}")
    _p(doc, "Người đại diện theo pháp luật: {{org_representative_name}}     Chức vụ: {{org_representative_title}}")
    _p(doc, "Số điện thoại: {{org_phone}}", after=2)

    _p(doc, "4. Phương thức nhận chi phí hỗ trợ mai táng:", bold=True)
    _p(doc, "{{bank_payment_check}} Tài khoản ngân hàng:", indent=0.5)
    _p(doc, "Tên chủ tài khoản: {{bank_account_name}}", indent=1)
    _p(doc, "Số tài khoản: {{bank_account_number}}", indent=1)
    _p(doc, "Ngân hàng: {{bank_name}}", indent=1)
    _p(doc, "{{cash_payment_check}} Tiền mặt", indent=0.5, after=3)
    _p(doc, "Tôi cam đoan những nội dung khai trên đây là đúng sự thật và chịu trách nhiệm trước pháp luật về nội dung khai của mình.", after=4)

    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    _no_borders(table)
    for cell in table.rows[0].cells:
        cell.width = Cm(8.7)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    headings = (
        "Ngày.... tháng.... năm ...\nNGƯỜI TIẾP NHẬN TỜ KHAI\n(Ký, ghi rõ họ tên)",
        "Ngày.... tháng.... năm ...\nNGƯỜI KHAI\n(Ký, ghi rõ họ tên)",
    )
    for cell, text in zip(table.rows[0].cells, headings):
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        for i, line in enumerate(text.split("\n")):
            r = p.add_run(line)
            _font(r, size=10.5, bold=i == 1, italic=i == 2)
            if i < 2:
                r.add_break()

    doc.save(OUTPUT)


if __name__ == "__main__":
    build()
