"""Tạo hai mẫu DOCX NQ40/NQ32 từ bố cục PDF do người dùng cung cấp."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_FOLDER = ROOT / "templates_word"


def _set_font(run, *, size: float = 12, bold: bool = False, italic: bool = False) -> None:
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def _format_paragraph(paragraph, *, after: float = 2, line: float = 1.0) -> None:
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    paragraph.paragraph_format.line_spacing = line


def _centered(document: Document, text: str, *, size: float = 12, bold: bool = False, italic: bool = False):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _format_paragraph(paragraph, after=1)
    _set_font(paragraph.add_run(text), size=size, bold=bold, italic=italic)
    return paragraph


def _field_line(document: Document, number: int, label: str, token: str, *, dots: int = 52) -> None:
    paragraph = document.add_paragraph()
    _format_paragraph(paragraph, after=6)
    _set_font(paragraph.add_run(f"{number}. {label}: "))
    _set_font(paragraph.add_run(f"{{{{{token}}}}}"), bold=True)
    _set_font(paragraph.add_run(" " + "." * dots))


def _remove_table_borders(table) -> None:
    properties = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "nil")
        borders.append(element)
    properties.append(borders)


def build_template(resolution_line: str, filename: str) -> Path:
    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.45)
    section.bottom_margin = Cm(1.35)
    section.left_margin = Cm(1.75)
    section.right_margin = Cm(1.75)

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(12)

    _centered(document, "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM", bold=True)
    motto = _centered(document, "Độc lập - Tự do - Hạnh phúc", size=13, bold=True)
    motto.runs[0].underline = True
    spacer = document.add_paragraph()
    _format_paragraph(spacer, after=2)

    _centered(document, "TỜ KHAI THÔNG TIN CÁ NHÂN", bold=True)
    _centered(document, resolution_line, size=11.5, bold=True)
    _centered(document, "CỦA HỘI ĐỒNG NHÂN DÂN THÀNH PHỐ HỒ CHÍ MINH", size=11.5, bold=True)

    recipient = document.add_paragraph()
    recipient.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _format_paragraph(recipient, after=10)
    _set_font(recipient.add_run("Kính gửi: Ủy ban nhân dân phường Minh Phụng"), size=13)

    _field_line(document, 1, "Họ và tên", "full_name")
    _field_line(document, 2, "Ngày, tháng, năm sinh", "date_of_birth")

    identity = document.add_paragraph()
    _format_paragraph(identity, after=6)
    _set_font(identity.add_run("3. CCCD số: "))
    _set_font(identity.add_run("{{citizen_id}}"), bold=True)
    _set_font(identity.add_run(", Ngày cấp: "))
    _set_font(identity.add_run("{{citizen_id_issue_date}}"), bold=True)
    _set_font(identity.add_run("; Nơi cấp: "))
    _set_font(identity.add_run("{{citizen_id_issue_place}}"), bold=True)

    _field_line(document, 4, "Hộ khẩu thường trú", "residence_address", dots=42)
    _field_line(document, 5, "Nơi tạm trú (nếu có)", "temporary_address", dots=43)
    _field_line(document, 6, "Số điện thoại", "phone_number", dots=55)
    _field_line(document, 7, "Nghề nghiệp", "occupation", dots=56)
    _field_line(document, 8, "Đơn vị công tác", "employer", dots=52)

    heading = document.add_paragraph()
    heading.paragraph_format.left_indent = Cm(1.45)
    _format_paragraph(heading, after=5)
    _set_font(heading.add_run("9. Thuộc đối tượng nhận hỗ trợ:"), size=13)

    options = (
        ("- Phụ nữ sinh đủ hai con trước 35 tuổi", "support_two_children_check"),
        (
            "- Phụ nữ mang thai và trẻ sơ sinh có thực hiện Sàng lọc trước sinh và "
            "Sàng lọc sơ sinh, cụ thể:",
            None,
        ),
        ("+ Hộ nghèo (theo tiêu chí của Thành phố), mã số: {{support_poor_detail}}", "support_poor_check"),
        ("+ Hộ cận nghèo (theo tiêu chí của Thành phố), mã số: {{support_near_poor_detail}}", "support_near_poor_check"),
        ("+ Đối tượng bảo trợ xã hội, cụ thể: {{support_social_detail}}", "support_social_check"),
        ("+ Đối tượng sống tại xã đảo: {{support_island_detail}}", "support_island_check"),
    )
    for index, (text, check_token) in enumerate(options):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.left_indent = Cm(1.45 if index < 2 else 2.2)
        _format_paragraph(paragraph, after=2)
        _set_font(paragraph.add_run(text))
        if check_token:
            _set_font(paragraph.add_run(f"  {{{{{check_token}}}}}"), size=13)

    declaration = document.add_paragraph()
    declaration.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    declaration.paragraph_format.first_line_indent = Cm(1.25)
    _format_paragraph(declaration, after=2)
    _set_font(
        declaration.add_run(
            "Tôi xin cam đoan chưa nhận hỗ trợ nội dung này tại những nơi khác và những lời "
            "khai trên đây của tôi là đúng sự thật, nếu có điều gì không đúng tôi xin chịu "
            "trách nhiệm trước pháp luật, đồng thời sẽ hoàn trả toàn bộ tiền hỗ trợ đã được nhận."
        )
    )
    request_line = document.add_paragraph()
    request_line.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    request_line.paragraph_format.first_line_indent = Cm(1.25)
    _format_paragraph(request_line, after=4)
    _set_font(request_line.add_run("Đề nghị Ủy ban nhân dân phường Minh Phụng chi hỗ trợ cho tôi."))

    signature = document.add_table(rows=1, cols=2)
    signature.autofit = False
    signature.columns[0].width = Cm(8.2)
    signature.columns[1].width = Cm(8.2)
    _remove_table_borders(signature)
    right = signature.cell(0, 1)
    right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    paragraphs = right.paragraphs
    signature_lines = (
        ("Ngày {{declaration_day}} tháng {{declaration_month}} năm {{declaration_year}}", False, True),
        ("Người khai", True, False),
        ("(ký và ghi rõ họ, tên)", False, True),
        ("", False, False),
        ("{{full_name}}", True, False),
    )
    for index, (text, bold, italic) in enumerate(signature_lines):
        paragraph = paragraphs[0] if index == 0 else right.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _format_paragraph(paragraph, after=1)
        _set_font(paragraph.add_run(text), bold=bold, italic=italic)

    properties = document.core_properties
    properties.title = resolution_line
    properties.subject = "Tờ khai thông tin cá nhân nhận hỗ trợ"
    properties.author = "Auto Form"

    OUTPUT_FOLDER.mkdir(parents=True, exist_ok=True)
    destination = OUTPUT_FOLDER / filename
    document.save(destination)
    return destination


if __name__ == "__main__":
    paths = (
        build_template(
            "NHẬN HỖ TRỢ THEO NGHỊ QUYẾT SỐ 40/NQ-HĐND NGÀY 11 THÁNG 12 NĂM 2024",
            "ho_tro_nq40.docx",
        ),
        build_template(
            "NHẬN HỖ TRỢ THEO NGHỊ QUYẾT SỐ 32/2025/NQ-HĐND NGÀY 28 THÁNG 5 NĂM 2025",
            "ho_tro_nq32.docx",
        ),
    )
    for path in paths:
        print(path)
