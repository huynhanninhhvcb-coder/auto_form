"""Kiểm thử luồng đầy đủ: quét CCCD người khai + Trích lục khai tử người mất
trong cùng một lượt tải lên, rồi tạo đơn hỏa táng/mai táng từ kết quả gộp.

Đây là kịch bản thực tế đơn hỗ trợ hỏa táng/mai táng luôn cần: thông tin
người đề nghị lấy từ CCCD của họ, thông tin người mất lấy từ Trích lục khai
tử — hai tài liệu, hai người khác nhau, trong cùng một lượt quét.

Dữ liệu dưới đây hoàn toàn giả định, không dùng tên/số của người dùng thực tế
trong bất kỳ ảnh nào được tải lên ứng dụng.
"""

from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from docx import Document

from services.batch_extraction_service import merge_extraction_results
from services.docx_service import generate_document
from services.extraction_service import extract_personal_information
from services.template_service import get_template


APPLICANT_CCCD_TEXT = """--- CCCD CHI TIẾT ---
Họ và tên / Full name:
NGUYỄN THỊ MAI
Ngày sinh: 10/05/1958
Giới tính: Nữ
--- CCCD BỐ CỤC ---
CĂN CƯỚC CÔNG DÂN
Citizen Identity Card
Số: 079158000777
--- CCCD ĐỊA CHỈ ---
Nơi thường trú: 12/3 Đường Số 5, Phường Minh Phụng, TP. Hồ Chí Minh
"""

DECEASED_EXTRACT_TEXT = """TRÍCH LỤC KHAI TỬ
(BẢN SAO)

Họ, chữ đệm, tên: LÊ VĂN HÙNG
Ngày, tháng, năm sinh: 20/03/1935
Giới tính: Nam    Dân tộc: Kinh    Quốc tịch: Việt Nam
Số định danh cá nhân: 079135000888
Đã chết vào lúc 06 giờ 30 phút, ngày 10/09/2026 ghi bằng chữ: Sáu giờ, ba mươi phút
Nơi chết: Bệnh viện Nhân dân 115, Thành phố Hồ Chí Minh
Đã được đăng ký khai tử tại: Ủy ban nhân dân phường Minh Phụng, Thành phố Hồ Chí
Minh
Số: 512/2026 ngày 11 tháng 9 năm 2026
"""


def _full_text(document: Document) -> str:
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.extend(paragraph.text for paragraph in cell.paragraphs)
    return "\n".join(parts)


class DeceasedBenefitWorkflowTests(unittest.TestCase):
    def setUp(self):
        applicant_result = extract_personal_information(
            APPLICANT_CCCD_TEXT, page_texts=[APPLICANT_CCCD_TEXT]
        )
        deceased_result = extract_personal_information(
            DECEASED_EXTRACT_TEXT, page_texts=[DECEASED_EXTRACT_TEXT]
        )
        self.merged = merge_extraction_results([applicant_result, deceased_result])

    def test_batch_extraction_separates_applicant_from_deceased(self):
        fields = self.merged.result.fields
        self.assertEqual(fields["full_name"], "NGUYỄN THỊ MAI")
        self.assertEqual(fields["citizen_id"], "079158000777")
        self.assertIn("Minh Phụng", fields["residence_address"])
        self.assertEqual(fields["deceased_full_name"], "LÊ VĂN HÙNG")
        self.assertEqual(fields["deceased_date_of_birth"], "20/03/1935")
        self.assertEqual(fields["deceased_citizen_id"], "079135000888")
        self.assertEqual(fields["deceased_death_date"], "10/09/2026")
        self.assertEqual(fields["death_certificate_number"], "512/2026")
        # Không có cảnh báo "người khác": đây là hai người hợp lệ trên hai
        # loại giấy tờ khác nhau, không phải dữ liệu mâu thuẫn của cùng một người.
        self.assertFalse(any("người khác" in warning for warning in self.merged.result.warnings))

    def _generate(self, template_id: str) -> str:
        with tempfile.TemporaryDirectory() as temporary_directory:
            output = Path(temporary_directory) / "filled.docx"
            generate_document(get_template(template_id), self.merged.result.fields, output)
            return _full_text(Document(output))

    def test_hoa_tang_don_contains_both_applicant_and_deceased_info(self):
        text = self._generate("ho_tro_hoa_tang")
        self.assertIn("NGUYỄN THỊ MAI", text)
        self.assertIn("079158000777", text)
        self.assertIn("LÊ VĂN HÙNG", text)
        self.assertIn("10/09/2026", text)
        self.assertNotIn("{{", text)

    def test_mai_tang_don_contains_both_applicant_and_deceased_info(self):
        text = self._generate("ho_tro_mai_tang")
        self.assertIn("NGUYỄN THỊ MAI", text)
        self.assertIn("079158000777", text)
        self.assertIn("LÊ VĂN HÙNG", text)
        self.assertIn("079135000888", text)
        self.assertIn("Nam", text)  # giới tính người mất
        self.assertNotIn("{{", text)


if __name__ == "__main__":
    unittest.main()
