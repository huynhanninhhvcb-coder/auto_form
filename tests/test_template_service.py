from __future__ import annotations

import unittest
from zipfile import ZipFile

from docx import Document

from services.template_service import get_template, list_templates


class TemplateServiceTests(unittest.TestCase):
    def test_lists_all_six_templates(self):
        ids = {item["id"] for item in list_templates()}
        self.assertEqual(
            ids,
            {
                "tro_cap_huu_tri",
                "ho_tro_hoa_tang",
                "ho_tro_mai_tang",
                "ho_tro_nq40",
                "ho_tro_nq32",
                "xac_dinh_khuyet_tat",
            },
        )

    def test_deceased_benefit_templates_share_a_common_base_plus_their_own_field(self):
        hoa_tang = get_template("ho_tro_hoa_tang")
        mai_tang = get_template("ho_tro_mai_tang")
        shared_base = ("full_name", "date_of_birth", "citizen_id", "residence_address", "deceased_full_name", "deceased_death_date")
        self.assertEqual(hoa_tang.required_fields[:6], shared_base)
        self.assertEqual(mai_tang.required_fields[:6], shared_base)
        # Hỏa táng cần thêm đối tượng hỏa táng; mai táng cần thêm nội dung đề
        # nghị và phương thức nhận hỗ trợ - hai mẫu không còn giống hệt nhau.
        self.assertEqual(hoa_tang.required_fields[6:], ("cremation_support_category",))
        self.assertEqual(mai_tang.required_fields[6:], ("request_content", "payment_method"))

    def test_new_template_word_files_exist_on_disk(self):
        self.assertTrue(get_template("ho_tro_hoa_tang").path.is_file())
        self.assertTrue(get_template("ho_tro_mai_tang").path.is_file())
        self.assertTrue(get_template("ho_tro_nq40").path.is_file())
        self.assertTrue(get_template("ho_tro_nq32").path.is_file())
        self.assertTrue(get_template("xac_dinh_khuyet_tat").path.is_file())

    def test_disability_determination_template_requires_procedure_type(self):
        template = get_template("xac_dinh_khuyet_tat")
        self.assertEqual(
            template.required_fields,
            ("full_name", "date_of_birth", "citizen_id", "residence_address", "disability_procedure_type"),
        )

    def test_nq_templates_require_identity_issue_and_support_group(self):
        expected = (
            "full_name",
            "date_of_birth",
            "citizen_id",
            "citizen_id_issue_date",
            "citizen_id_issue_place",
            "residence_address",
            "support_category",
        )
        self.assertEqual(get_template("ho_tro_nq40").required_fields, expected)
        self.assertEqual(get_template("ho_tro_nq32").required_fields, expected)

    def test_nq_templates_keep_official_blank_signature_and_full_width_dot_leaders(self):
        for template_id in ("ho_tro_nq40", "ho_tro_nq32"):
            template = get_template(template_id)
            document = Document(template.path)
            text = "\n".join(paragraph.text for paragraph in document.paragraphs)
            table_text = "\n".join(
                cell.text for table in document.tables for row in table.rows for cell in row.cells
            )
            with ZipFile(template.path) as archive:
                xml = archive.read("word/document.xml")
            with self.subTest(template_id=template_id):
                self.assertEqual(text.count("{{full_name}}"), 1)
                self.assertIn("Ngày..... tháng.... năm 20.....", table_text)
                self.assertGreaterEqual(xml.count(b'w:leader="dot"'), 7)

    def test_unknown_template_returns_none(self):
        self.assertIsNone(get_template("khong-ton-tai"))


if __name__ == "__main__":
    unittest.main()
