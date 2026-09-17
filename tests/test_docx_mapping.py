from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from docx import Document

from services.docx_service import generate_document
from services.template_service import get_template


def _paragraph_containing(document: Document, text: str) -> str:
    for paragraph in document.paragraphs:
        if text in paragraph.text:
            return paragraph.text
    raise AssertionError(f"Không tìm thấy paragraph chứa {text!r}")


def _line_containing(full_text: str, text: str) -> str:
    for line in full_text.splitlines():
        if text in line:
            return line
    raise AssertionError(f"Không tìm thấy dòng chứa {text!r}")


class RetirementAssistanceDocxMappingTests(unittest.TestCase):
    def test_fills_multi_field_lines_at_their_own_labels(self):
        with tempfile.TemporaryDirectory() as temporary_directory:
            output = Path(temporary_directory) / "filled.docx"
            generate_document(
                get_template("tro_cap_huu_tri"),
                {
                    "full_name": "NGUYỄN VĂN AN",
                    "date_of_birth": "1-2-1950",
                    "gender": "Nam",
                    "ethnic_group": "Kinh",
                    "citizen_id": "079 150 001 234",
                    "residence_address": "Phường Minh Phụng, TP. Hồ Chí Minh",
                    "contact_address": "Phường Minh Phụng, TP. Hồ Chí Minh",
                    "bank_account_name": "NGUYỄN VĂN AN",
                    "bank_account_number": "0123456789",
                    "bank_name": "Ngân hàng Kiểm Thử",
                },
                output,
            )

            document = Document(output)

        name_line = _paragraph_containing(document, "1. Họ, chữ đệm, tên")
        identity_line = _paragraph_containing(document, "2. Ngày, tháng, năm sinh")
        citizen_id_line = _paragraph_containing(document, "3. Thẻ Căn cước")
        bank_name_line = _paragraph_containing(document, "- Tên tài khoản")
        bank_number_line = _paragraph_containing(document, "- Số tài khoản")

        self.assertRegex(name_line, r":\s+NGUYỄN VĂN AN$")
        self.assertRegex(citizen_id_line, r":\s+079150001234$")
        self.assertRegex(bank_name_line, r":\s+NGUYỄN VĂN AN$")

        self.assertEqual(identity_line.count("Giới tính"), 1)
        self.assertEqual(identity_line.count("Dân tộc"), 1)
        self.assertLess(identity_line.index("01/02/1950"), identity_line.index("Giới tính"))
        self.assertLess(identity_line.index("Nam"), identity_line.index("Dân tộc"))
        self.assertLess(identity_line.index("Dân tộc"), identity_line.index("Kinh"))

        self.assertEqual(bank_number_line.count("Ngân hàng:"), 1)
        self.assertLess(bank_number_line.index("0123456789"), bank_number_line.index("Ngân hàng"))
        self.assertLess(bank_number_line.index("Ngân hàng"), bank_number_line.index("Ngân hàng Kiểm Thử"))

    def test_keeps_empty_optional_bank_fields_blank(self):
        with tempfile.TemporaryDirectory() as temporary_directory:
            output = Path(temporary_directory) / "filled.docx"
            generate_document(
                get_template("tro_cap_huu_tri"),
                {
                    "full_name": "NGUYỄN VĂN AN",
                    "date_of_birth": "01/02/1950",
                    "citizen_id": "079150001234",
                    "residence_address": "Phường Minh Phụng",
                },
                output,
            )

            document = Document(output)

        bank_name_line = _paragraph_containing(document, "- Tên tài khoản")
        bank_number_line = _paragraph_containing(document, "- Số tài khoản")

        self.assertNotIn("NGUYỄN VĂN AN", bank_name_line)
        self.assertNotIn("079150001234", bank_number_line)
        self.assertEqual(bank_number_line.count("Ngân hàng"), 1)


def _full_text(document: Document) -> str:
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.extend(paragraph.text for paragraph in cell.paragraphs)
    return "\n".join(parts)


class DeceasedBenefitDocxMappingTests(unittest.TestCase):
    """Hai mẫu mới dùng {{placeholder}} thuần túy, không cần hàm map riêng."""

    def _generate(self, template_id: str, fields: dict) -> str:
        with tempfile.TemporaryDirectory() as temporary_directory:
            output = Path(temporary_directory) / "filled.docx"
            generate_document(get_template(template_id), fields, output)
            return _full_text(Document(output))

    def test_hoa_tang_template_fills_applicant_and_deceased_fields(self):
        text = self._generate(
            "ho_tro_hoa_tang",
            {
                "full_name": "NGUYỄN VĂN AN",
                "date_of_birth": "1-2-1980",
                "citizen_id": "079 080 000123",
                "residence_address": "123 Lý Nam Đế",
                "deceased_relationship": "Con ruột",
                "deceased_full_name": "NGUYỄN VĂN BA",
                "deceased_death_date": "2-3-2025",
            },
        )
        self.assertIn("NGUYỄN VĂN AN", text)
        self.assertIn("01/02/1980", text)
        self.assertIn("079080000123", text)
        self.assertIn("NGUYỄN VĂN BA", text)
        self.assertIn("02/03/2025", text)
        self.assertNotIn("{{", text)

    def test_mai_tang_template_fills_deceased_details_and_leaves_org_blank(self):
        text = self._generate(
            "ho_tro_mai_tang",
            {
                "full_name": "NGUYỄN VĂN AN",
                "citizen_id": "079080000123",
                "residence_address": "123 Lý Nam Đế",
                "deceased_full_name": "NGUYỄN VĂN BA",
                "deceased_death_date": "02/03/2025",
                "deceased_citizen_id": "079 050 000456",
                "deceased_gender": "Nam",
            },
        )
        self.assertIn("079050000456", text)
        self.assertIn("Nam", text)
        # Không điền tổ chức -> vẫn phải để trống, không lộ placeholder thô.
        self.assertNotIn("{{", text)

    def test_no_leftover_placeholder_when_field_missing_entirely(self):
        text = self._generate(
            "ho_tro_hoa_tang",
            {"full_name": "NGUYỄN VĂN AN"},
        )
        self.assertNotIn("{{", text)

    def test_hoa_tang_checks_exactly_the_selected_cremation_category(self):
        text = self._generate(
            "ho_tro_hoa_tang",
            {
                "full_name": "NGUYỄN VĂN AN",
                "cremation_support_category": "Hộ cận nghèo",
                "cremation_support_detail": "CN-456",
            },
        )
        self.assertNotIn("{{", text)
        near_poor_line = _line_containing(text, "Hộ cận nghèo")
        self.assertIn("☒", near_poor_line)
        self.assertIn("CN-456", near_poor_line)
        # Đối tượng hộ nghèo không được chọn nên vẫn phải bỏ trống mã số của nó.
        poor_line = _line_containing(text, "Hộ nghèo (")
        self.assertIn("☐", poor_line)
        self.assertNotIn("CN-456", poor_line)
        # Toàn bộ 19 dòng đối tượng khác phải giữ nguyên "☐", chỉ một dòng "☒".
        self.assertEqual(text.count("☒"), 1)

    def test_mai_tang_checks_the_selected_payment_method_and_fills_death_time(self):
        text = self._generate(
            "ho_tro_mai_tang",
            {
                "full_name": "NGUYỄN VĂN AN",
                "deceased_full_name": "NGUYỄN VĂN BA",
                "deceased_death_date": "02/03/2025",
                "deceased_death_hour": "14",
                "deceased_death_minute": "5",
                "gender": "Nam",
                "payment_method": "Tiền mặt",
                "request_content": "Đề nghị hỗ trợ chi phí mai táng",
            },
        )
        self.assertNotIn("{{", text)
        cash_line = _line_containing(text, "Tiền mặt")
        self.assertTrue(cash_line.startswith("☒"))
        bank_line = _line_containing(text, "Tài khoản ngân hàng")
        self.assertTrue(bank_line.startswith("☐"))
        self.assertIn("14 giờ 5 phút", text)
        self.assertIn("Đề nghị hỗ trợ chi phí mai táng", text)
        self.assertLess(text.index("Nội dung đề nghị"), text.index("2. Thông tin người chết"))
        self.assertIn("Nam/Nữ: Nam", text)


class NqSupportDocxMappingTests(unittest.TestCase):
    def test_both_nq_templates_fill_fields_and_select_exactly_one_category(self):
        fields = {
            "full_name": "NGUYỄN THỊ MAI",
            "date_of_birth": "1-2-1990",
            "citizen_id": "079 090 001 234",
            "citizen_id_issue_date": "3-4-2021",
            "citizen_id_issue_place": "Cục Cảnh sát QLHC về TTXH",
            "residence_address": "123 Lý Nam Đế, Phường Minh Phụng",
            "temporary_address": "Như trên",
            "phone_number": "0912 345 678",
            "occupation": "Nhân viên",
            "employer": "Công ty A",
            "support_category": "Hộ cận nghèo",
            "support_detail": "CN-123",
        }
        for template_id, resolution in (
            ("ho_tro_nq40", "40/NQ-HĐND"),
            ("ho_tro_nq32", "32/2025/NQ-HĐND"),
        ):
            with self.subTest(template_id=template_id), tempfile.TemporaryDirectory() as temporary_directory:
                output = Path(temporary_directory) / "filled.docx"
                generate_document(get_template(template_id), fields, output)
                text = _full_text(Document(output))

            self.assertIn(resolution, text)
            self.assertIn("NGUYỄN THỊ MAI", text)
            self.assertIn("079090001234", text)
            self.assertIn("03/04/2021", text)
            self.assertIn("CN-123", text)
            self.assertEqual(text.count("☒"), 1)
            self.assertEqual(text.count("☐"), 4)
            self.assertNotIn("{{", text)


if __name__ == "__main__":
    unittest.main()
